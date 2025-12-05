from docx.shared import Pt
from docx.oxml import parse_xml
from document.format.helpers import format_number, format_decimal_2

def create_table_with_merged_cells(sheet, docx_table, start_row, num_rows, start_col, num_cols, key, headers):
    """
    Populates a docx table from an Excel sheet, handling merged cells.

    Args:
        sheet: The openpyxl worksheet.
        docx_table: The python-docx table object.
        start_row (int): The starting row index in the Excel sheet.
        num_rows (int): The number of rows in the table.
        start_col (int): The starting column index in the Excel sheet.
        num_cols (int): The number of columns in the table.
        key (str): The language key ('es' or 'en').
        headers (list): The list of header strings.
    """

    merged_cells = {str(mc) for mc in sheet.merged_cells.ranges}

    for i in range(num_rows):
        for j in range(num_cols):
            excel_cell = sheet.cell(row=start_row + i, column=start_col + j)
            docx_cell = docx_table.cell(i, j)

            val = excel_cell.value
            text = str(val) if val is not None else ""

            if isinstance(val, (int, float)):
                header_text = headers[j]
                if header_text in ["Desplazamiento (mm)", "Displacement (mm)", "Desplazamiento Relativo (mm)", "Relative Displacement (mm)", "Desplazamiento Permisible (mm)", "Allowable Displacement (mm)"]:
                    text = format_decimal_2(val)
                else:
                    text = format_number(val)

            original_text_for_coloring = text
            header_text = headers[j]
            if header_text in ["Tipo", "Type"]:
                if key == 'es':
                    if text == 'PM': text = 'MF'
                    elif text == 'AM': text = 'EA'
                elif key == 'en':
                    if text == 'AM': text = 'AE'
            
            if header_text in ["Verificación", "Verification"]:
                if key == 'en':
                    if text.lower().strip() in ['si', 'sí']: text = 'Yes'
                    elif text.lower().strip() == 'no': text = 'No'
                
                if original_text_for_coloring.lower().strip() in ["si", "sí", "yes", "true", "1"]:
                    shade = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="C6EFCE"/>')
                    docx_cell._tc.get_or_add_tcPr().append(shade)
                elif original_text_for_coloring.lower().strip() in ["no", "false", "0"]:
                    shade = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="FFC7CE"/>')
                    docx_cell._tc.get_or_add_tcPr().append(shade)

            docx_cell.text = text
            p = docx_cell.paragraphs[0]
            p.alignment = 1
            if not p.runs:
                p.add_run()
            run = p.runs[0]
            run.font.name = "Arial"
            if i == 0: # Header row
                run.bold = True
                shade = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="D9D9D9"/>')
                docx_cell._tc.get_or_add_tcPr().append(shade)


    # Apply merges
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        
        # Check if the merged range is within the bounds of our table
        if min_row >= start_row and max_row <= start_row + num_rows and min_col >= start_col and max_col <= start_col + num_cols:
            
            # Adjust to 0-based index for docx table
            start_merge_row = min_row - start_row
            start_merge_col = min_col - start_col
            end_merge_row = max_row - start_row
            end_merge_col = max_col - start_col

            top_left_cell = docx_table.cell(start_merge_row, start_merge_col)
            bottom_right_cell = docx_table.cell(end_merge_row, end_merge_col)
            
            top_left_cell.merge(bottom_right_cell)